"""验证输入边界、引用可信边界和网络失败；真实 API 需另行手动验收。"""
import json
import unittest
from io import BytesIO
from unittest.mock import Mock, patch
import requests
from PIL import Image
from docx import Document
from studylens.documents import load_documents, parse_document
from studylens.retrieval import Retriever
from studylens.llm import ModelConfig, ModelError, generate_answer, validate_answer
from studylens.ocr import extract_image_text
from studylens.office import extract_docx_text, extract_pdf_text


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.chunks = load_documents([("笔记.md", "# 学习\n## 苹果\n苹果是这里的检索目标。\n## 香蕉\n香蕉在第二节。\n".encode("utf-8"))])
        self.index = Retriever(self.chunks)
        self.hits = self.index.search("苹果", min_score=0)

    def test_heading_and_lines(self):
        apple = next(c for c in self.chunks if c.section == "苹果")
        self.assertEqual((apple.start_line, apple.end_line), (2, 3))
        self.assertIn("苹果是", apple.text)

    def test_long_chunk_preserves_content(self):
        text = "一二三四五六七八九十" * 80
        chunks = parse_document("a.txt", text.encode("utf-8"), "D1", 100, 20)
        rebuilt = chunks[0].text + "".join(c.text[20:] for c in chunks[1:])
        self.assertEqual(rebuilt, text)

    def test_invalid_documents(self):
        for name, data in [("a.md", b""), ("a.md", b"\xff"), ("a.pdf", b"abc"),
                           ("a.txt", b"a" * (512 * 1024 + 1))]:
            with self.subTest(name=name, size=len(data)):
                with self.assertRaises(ValueError):
                    load_documents([(name, data)])

    def test_image_ocr_result_can_enter_document_pipeline(self):
        image = BytesIO()
        Image.new("RGB", (20, 20), "white").save(image, format="PNG")
        engine = Mock(return_value=Mock(txts=("极化码", "列表译码"), scores=(0.9, 0.8)))
        recognized = extract_image_text("课堂笔记.png", image.getvalue(), engine)
        chunks = load_documents([("课堂笔记.png.ocr.txt", recognized.text.encode("utf-8"))])
        self.assertEqual(recognized.line_count, 2)
        self.assertAlmostEqual(recognized.confidence, 0.85)
        self.assertIn("列表译码", chunks[0].text)

    def test_image_ocr_rejects_bad_or_empty_results(self):
        with self.assertRaisesRegex(ValueError, "有效图片"):
            extract_image_text("坏图.png", b"not-an-image", Mock())
        image = BytesIO()
        Image.new("RGB", (20, 20), "white").save(image, format="PNG")
        with self.assertRaisesRegex(ValueError, "未识别出文字"):
            extract_image_text("空白.png", image.getvalue(), Mock(return_value=Mock(txts=(), scores=())))

    def test_docx_extracts_headings_paragraphs_and_tables(self):
        data = BytesIO()
        document = Document()
        document.add_heading("极化码", level=1)
        document.add_paragraph("SC 是逐次消除译码。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "算法"
        table.cell(0, 1).text = "SCL"
        document.save(data)
        extracted = extract_docx_text("课程笔记.docx", data.getvalue())
        self.assertIn("# 极化码", extracted.text)
        self.assertIn("算法 | SCL", extracted.text)

    def test_scanned_pdf_uses_ocr_and_preserves_page_heading(self):
        pdf = BytesIO()
        Image.new("RGB", (100, 100), "white").save(pdf, format="PDF")
        engine = Mock(return_value=Mock(txts=("扫描笔记内容",), scores=(0.9,)))
        extracted = extract_pdf_text("扫描笔记.pdf", pdf.getvalue(), Mock(return_value=engine))
        self.assertIn("# 第 1 页", extracted.text)
        self.assertIn("扫描笔记内容", extracted.text)
        self.assertEqual(extracted.ocr_pages, (1,))

    @patch("studylens.office.PdfReader")
    def test_text_pdf_does_not_start_ocr(self, reader_class):
        reader_class.return_value = Mock(
            is_encrypted=False,
            pages=[Mock(extract_text=Mock(return_value="SC 是逐次消除译码。"))],
        )
        engine_factory = Mock()
        extracted = extract_pdf_text("电子笔记.pdf", b"valid-enough-for-mock", engine_factory)
        self.assertIn("SC 是逐次消除译码", extracted.text)
        self.assertEqual(extracted.ocr_pages, ())
        engine_factory.assert_not_called()

    def test_office_files_reject_invalid_content(self):
        with self.assertRaisesRegex(ValueError, "有效的 DOCX"):
            extract_docx_text("坏文档.docx", b"bad")
        with self.assertRaisesRegex(ValueError, "有效的 PDF"):
            extract_pdf_text("坏文档.pdf", b"bad", Mock())

    def test_bom_and_duplicate_names(self):
        chunks = load_documents([("a.txt", "\ufeff苹果".encode("utf-8")), ("a.txt", "香蕉".encode("utf-8"))])
        self.assertEqual(len({c.id for c in chunks}), 2)
        self.assertNotIn("\ufeff", chunks[0].text)

    def test_retrieval_and_empty_question(self):
        self.assertEqual(self.hits[0].chunk.section, "苹果")
        self.assertEqual(self.index.search(" "), [])
        self.assertEqual(self.index.search("XYZ"), [])

    def test_unknown_citation_rejected(self):
        with self.assertRaises(ModelError):
            validate_answer({"answer": "a", "citations": ["fake"], "insufficient": False}, {"D1-1"})

    def test_wrong_schema_rejected(self):
        for result in [[], {"answer": "a", "citations": "D1-1", "insufficient": False},
                       {"answer": "a", "citations": [], "insufficient": False},
                       {"answer": "a", "citations": [], "insufficient": "false"}]:
            with self.assertRaises(ModelError):
                validate_answer(result, {"D1-1"})

    def test_insufficient_overrides_unsupported_answer(self):
        result = validate_answer({"answer": "编造的答案", "citations": [], "insufficient": True}, set())
        self.assertIn("资料不足", result["answer"])

    @patch("studylens.llm.requests.post")
    def test_no_evidence_never_calls_network(self, post):
        result = generate_answer("问题", [], ModelConfig("", "", ""))
        self.assertTrue(result["insufficient"])
        post.assert_not_called()

    @patch("studylens.llm.requests.post")
    def test_mock_api_success(self, post):
        identifier = self.hits[0].chunk.id
        content = json.dumps({"answer": "苹果是目标。", "citations": [identifier], "insufficient": False})
        post.return_value = Mock(status_code=200)
        post.return_value.json.return_value = {"choices": [{"message": {"content": content}, "finish_reason": "stop"}]}
        answer = generate_answer("苹果", self.hits, ModelConfig("test-only", "https://example.invalid/v1", "test"))
        self.assertEqual(answer["citations"], [identifier])
        self.assertFalse(post.call_args.kwargs["allow_redirects"])
        self.assertEqual(post.call_args.args[0], "https://example.invalid/v1/chat/completions")

    @patch("studylens.llm.requests.post")
    def test_http_and_timeout_errors(self, post):
        config = ModelConfig("test-only", "https://example.invalid", "test")
        post.return_value = Mock(status_code=401)
        with self.assertRaisesRegex(ModelError, "401"):
            generate_answer("苹果", self.hits, config)
        post.side_effect = requests.Timeout()
        with self.assertRaisesRegex(ModelError, "超时"):
            generate_answer("苹果", self.hits, config)

    @patch("studylens.llm.requests.post")
    def test_malformed_and_truncated_responses(self, post):
        config = ModelConfig("test-only", "https://example.invalid", "test")
        post.return_value = Mock(status_code=200)
        post.return_value.json.return_value = {"choices": [{"message": {"content": "not-json"}}]}
        with self.assertRaises(ModelError):
            generate_answer("苹果", self.hits, config)
        post.return_value.json.return_value = {"choices": [{"finish_reason": "length"}]}
        with self.assertRaisesRegex(ModelError, "截断"):
            generate_answer("苹果", self.hits, config)


if __name__ == "__main__":
    unittest.main()
